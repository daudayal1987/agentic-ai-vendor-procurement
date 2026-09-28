from io import BytesIO
from uuid import UUID

from docx import Document

from app.documents.parsing.interface import (
    DocumentParser,
    ParsedDocument,
    ParsedSegment,
)


class DOCXDocumentParser(DocumentParser):

    def supports(self, content_type: str) -> bool:
        return content_type == (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )

    def parse(
        self,
        document_id: UUID,
        content: bytes,
        content_type: str,
    ) -> ParsedDocument:

        document = Document(BytesIO(content))

        segments: list[ParsedSegment] = []

        for paragraph_index, paragraph in enumerate(
            document.paragraphs
        ):
            text = paragraph.text

            if not text.strip():
                continue

            segments.append(
                ParsedSegment(
                    text=text,
                    metadata={
                        "paragraph_index": paragraph_index,
                    },
                )
            )

        text = "\n\n".join(
            segment.text
            for segment in segments
        )

        return ParsedDocument(
            document_id=document_id,
            text=text,
            content_type=content_type,
            segments=tuple(segments),
        )